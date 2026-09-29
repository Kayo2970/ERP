"""
generate_manual_docx.py — Comprehensive Word Manual Generator for LEADS ERP.
Generates docs/LEADS_ERP_User_Manual.docx covering all 23 modules,
7-tier role privileges, operational workflows, and security architecture.
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_manual():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    
    # Title Page / Header Block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("LEADS NEXT GEN CENTRE\nOPERATIONS & PRIVILEGES MANUAL")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 60, 110)
    
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Institutional User Guide, Role Privilege Hierarchy & Complete 23-Module Standard Operating Procedures")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    
    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Platform:", "LEADS ERP All-in-One Operations Portal (Production v2.0)"),
        ("Institution:", "LEADS Next Gen Centre, M.S. Ramaiah University of Applied Sciences (MSRUAS)"),
        ("System Architecture:", "Next.js 16 + React 19 + AES-256-GCM Encrypted Flat-File Store"),
        ("Document Classification:", "Official Institutional Operations Manual"),
        ("Target Roles:", "Tiers 1 to 7 (Super Users, Faculty Heads, Committee Leads, Students, Alumni)"),
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.5)
        set_cell_background(c1, "F1F5F9")
        set_cell_background(c2, "FAFAFA")
        set_cell_margins(c1, 70, 70, 120, 120)
        set_cell_margins(c2, 70, 70, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(15, 60, 110)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(9.5)
        
    doc.add_paragraph()
    
    # Section 1
    doc.add_heading("1. Executive Platform Overview & Architecture", level=1)
    doc.add_paragraph(
        "The LEADS Next Gen ERP Dashboard is a unified operations portal developed for the LEADS Next Gen Centre at MSRUAS. "
        "The platform coordinates 23 distinct operational modules, spanning event planning, task delegation, digital gate passes, "
        "mobile wallet integration (Apple Wallet & Google Wallet), multi-tier financial reimbursement audits, equipment procurement, "
        "dynamic form builders, and student performance ratings."
    )
    doc.add_paragraph(
        "Key System Characteristics:\n"
        "• Zero External Database Dependency: All data resides in encrypted collections (data/*.json) protected with AES-256-GCM encryption.\n"
        "• Reactive Cross-Session Sync: Changes made by one user automatically propagate to other active sessions within 7 seconds via 'leads-data-sync' events.\n"
        "• Cryptographic Credentials: Passwords are protected via Node's native scrypt hashing with unique 16-byte salts and timing-safe comparisons."
    )
    
    # Section 2
    doc.add_heading("2. Role Privilege & Access Hierarchy (Tiers 1 to 7)", level=1)
    doc.add_paragraph(
        "Access privileges are strictly partitioned into 7 authority tiers combined with organizational divisions:"
    )
    
    tier_table = doc.add_table(rows=1, cols=4)
    tier_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tier_table.rows[0].cells
    hdr[0].width = Inches(0.9)
    hdr[1].width = Inches(1.5)
    hdr[2].width = Inches(1.3)
    hdr[3].width = Inches(3.0)
    
    headers = ["Tier", "Role Title", "Division", "Authorized System Scope"]
    for i, h in enumerate(headers):
        set_cell_background(hdr[i], "0F3C6E")
        set_cell_margins(hdr[i], 80, 80, 80, 80)
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        
    tiers_data = [
        ("Tier 1", "Super User", "Core Committee", "Root platform authority, audit logs, encrypted snapshots, system configuration, policy overrides, and live persona switching."),
        ("Tier 2", "Centre Head", "Faculty", "Campus-wide authority, final event approvals, Level-2 reimbursement clearance, major procurement approval, and VIP guest directory."),
        ("Tier 2.5", "GG Campus Head", "Faculty", "Campus-specific oversight for Gnanagangothri (GG) campus operations and task coordination."),
        ("Tier 3", "Department Head / Event Lead", "Faculty / Core", "Directs specific departments (Events, Media, Stage, Logistics). Proposal approvals, Level-1 reimbursement audits, task delegation."),
        ("Tier 4", "Advisory Board / Faculty Advisor", "Faculty", "Strategic read access to institutional performance analytics, committee ratings, and event documentation."),
        ("Tier 5", "Core Committee Lead", "Core Committee", "Event setup, sub-committee member allocations, pass studio badge issuance, form builder, and task planning."),
        ("Tier 6", "Training Associate / Student", "Training Assoc.", "Personal dashboard: task execution & progress updates, reimbursement claim submission, and digital visiting card generation."),
        ("Tier 7", "Alumni / Guest", "Alumni / Guest", "Read-only access to relevant past archives, digital visiting card exchange, and public event passes."),
    ]
    
    for row_data in tiers_data:
        row = tier_table.add_row().cells
        row[0].width = Inches(0.9)
        row[1].width = Inches(1.5)
        row[2].width = Inches(1.3)
        row[3].width = Inches(3.0)
        for i, val in enumerate(row_data):
            set_cell_background(row[i], "FFFFFF" if "Tier 1" in row_data[0] or "Tier 3" in row_data[0] or "Tier 5" in row_data[0] else "F8FAFC")
            set_cell_margins(row[i], 60, 60, 80, 80)
            p = row[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
                
    doc.add_paragraph()
    
    # Section 3
    doc.add_heading("3. Complete 23-Module Operational Guide", level=1)
    
    modules = [
        ("3.1 Home Dashboard",
         "The executive landing page upon logging in. Displays KPI stat cards (Active Events, Open Tasks, Expenditure vs. Budget), "
         "an Action Required inbox highlighting pending approvals, upcoming calendar milestones, and recent system audit feeds."),
        
        ("3.2 Events Management",
         "Coordinates event lifecycles: Proposal -> Approval -> Active Execution -> Completion -> Archival. Organizers can attach "
         "functional sub-committees (Stage, Logistics, Social Media, Hospitality), assign student leads, establish budgets, and monitor progress."),
        
        ("3.3 Event Passes & Live Scanner Kiosk",
         "Enables pass creation and gate check-in. The Pass Studio generates badges with unique cryptographic serials and QR codes. "
         "On event day, gate volunteers use the camera-based Check-in Scanner on mobile or laptop for instant audio-visual check-in validation."),
        
        ("3.4 Apple Wallet & Google Wallet Integration",
         "Delivers mobile digital passes. Attendees tap 'Add to Apple Wallet' to install .pkpass badges featuring event schedules and "
         "lock-screen notifications when near campus. Android users install passes to Google Wallet via one-click JWT links."),
        
        ("3.5 Tasks, Delegation & Gantt Timeline",
         "Assigns deliverables to individuals or entire sub-committees. Includes priority levels, file attachments, and a Gantt Timeline "
         "visualizing critical path dependencies. Automatically sends email warnings 24 hours and 2 hours before deadlines."),
        
        ("3.6 Performance Ratings & Committee Evaluations",
         "Conducts post-event reviews across standard institutional criteria: Timeliness, Deliverable Quality, Teamwork, and Initiative. "
         "Scores aggregate into student recognition records and departmental efficiency indexes."),
        
        ("3.7 Procurement & Equipment Requisitions",
         "Coordinates resource sourcing (printing, electronics, stage equipment). Organizers upload vendor quotation estimates. "
         "Department Heads approve standard requests; high-value purchases route to the Centre Head for final clearance."),
        
        ("3.8 Financial Reimbursements (Dual-Gate Audit)",
         "Student organizers submit claims with receipt photos. The dual-gate pipeline enforces: Gate 1 (Faculty Head verification) "
         "followed by Gate 2 (Centre Head clearance). Disbursed claims automatically debit linked event budgets."),
        
        ("3.9 Budgeting, P&L & Income Sources",
         "Maintains departmental balance sheets. Tracks university allocations, external sponsorships, and ticket revenues with "
         "real-time calculation of Budget vs. Actual expenditure variances."),
        
        ("3.10 Design Portal & OCR Proofreading Engine",
         "Manages promotional posters and social media creatives. The built-in AI OCR engine scans uploaded graphics to detect typos, "
         "mismatched dates, or spelling errors before designated Media Faculty provide final sign-off."),
        
        ("3.11 Dynamic Form Builder & Public Submissions",
         "Builds custom RSVP, registration, and feedback surveys with custom public URLs (e.g., /forms/symposium-rsvp). "
         "Generates high-resolution QR codes and enables response exports to Word (.docx) and CSV."),
        
        ("3.12 Digital Visiting Cards & 3D Interactive Keycard",
         "Every verified member receives a digital business card (/card/[slug]). Features an interactive 3D WebGL badge, "
         "one-tap vCard phone address book download, and Apple/Google Wallet pass installation."),
        
        ("3.13 Guest Directory & VIP Invitation Engine",
         "Maintains records of guest lecturers, VIP dignitaries, and alumni. Features business card photo OCR to auto-populate "
         "contact profiles and one-click dispatch of formal personalized invitation emails."),
        
        ("3.14 Announcements & Scoped Broadcasts",
         "Broadcasts notices scoped to All Members, Core Committee Only, or Faculty Only. Pinned alerts appear on the dashboard "
         "and trigger automated outbound notification emails."),
        
        ("3.15 Master Calendar & Festival Schedule",
         "Unified institutional timetable displaying event dates, milestone deadlines, and university academic holidays across campuses. "
         "Supports one-click export to Google and Apple Calendar via iCal."),
        
        ("3.16 Executive Event Reports & PDF/Word Exports",
         "Post-event documentation engine compiling executive summaries, attendance analytics, final financial balance sheets, "
         "and photo galleries into formatted PDF and Word (.docx) reports for university management."),
        
        ("3.17 Unified Approvals Inbox",
         "A consolidated decision-making inbox for Tier 2/3 heads. Aggregates pending event proposals, procurement requests, "
         "reimbursement claims, and creative assets into a single list with inline approval actions."),
        
        ("3.18 Members Directory & User Management",
         "Administers member accounts, promotions, division assignments, and offboarding. Generates cryptographic activation links "
         "for new inductees and supports forced password resets."),
        
        ("3.19 Custom Group Policies & Access Thresholds",
         "Enables granular access control customization without code modifications, such as granting a specific student lead "
         "temporary access to financial budgeting modules."),
        
        ("3.20 Email Delivery Engine & Queue Logs",
         "Monitors outbound email health (task reminders, birthday greetings, pass deliveries, and circulars) with real-time "
         "delivery receipts, bounce tracking, and SMTP test tools."),
        
        ("3.21 System & Security Settings",
         "Configures institutional portal branding, SMTP mail relays (Google Workspace, Office 365, or local Postfix), and "
         "Apple/Google digital wallet credentials."),
        
        ("3.22 Encrypted Backup & Restore",
         "Guarantees disaster recovery. Generates master-key encrypted .zip snapshot archives of all 28 database collections "
         "and uploaded files for offline institutional storage and point-in-time restoration."),
        
        ("3.23 Mobile PWA Installation",
         "Enables users to install LEADS ERP onto Apple iOS or Android devices as a native-like application with offline asset caching "
         "and full-screen viewing.")
    ]
    
    for title, desc in modules:
        doc.add_heading(title, level=2)
        doc.add_paragraph(desc)
        
    doc.add_heading("4. Role-Specific Operational Workflows", level=1)
    
    workflows = [
        ("4.1 Student Members & Training Associates (Tiers 5-6)",
         "• Daily Routine: Log in to review assigned tasks on the Home Dashboard.\n"
         "• Task Completion: Open the task, update notes, attach deliverable files, and click 'Submit for Review'.\n"
         "• Expense Claim: After purchasing event items, take a photo of the bill, open Reimbursements, click '+ Submit Claim', and enter details.\n"
         "• Networking: Share your personal QR code from Visiting Card or tap 'Save Contact' to exchange vCard details."),
        
        ("4.2 Event Organizers & Committee Leads (Tiers 3, 5)",
         "• Planning: Create event proposal, define sub-committees (Stage, Logistics, Media), and assign student volunteers.\n"
         "• Ticketing: In Event Passes -> Studio, issue attendee passes and dispatch automated email confirmations.\n"
         "• Gate Check-in: Assign volunteers to operate the Check-in Scanner on mobile devices on event day.\n"
         "• Closure: Ensure tasks are signed off, compile the Event Report, and score committee ratings."),
        
        ("4.3 Faculty Advisors & Department Heads (Tiers 2-4)",
         "• Approvals: Check the Approvals inbox daily to clear event proposals and procurement requisitions.\n"
         "• Reimbursement Audit: Review Gate-1 reimbursement claims, verifying bill scans against amounts claimed before approving.\n"
         "• Quality Oversight: Inspect creative designs in the Design Portal and review post-event analytical reports."),
        
        ("4.4 Centre Head & Executive Leadership (Tier 2)",
         "• Financial Oversight: Approve final Level-2 reimbursement payouts and major institutional capital expenditures.\n"
         "• Strategic Planning: Review the master P&L budget sheet to allocate funding across upcoming academic semesters.\n"
         "• VIP Engagements: Direct high-level dignitary invitations via the Guest Directory."),
        
        ("4.5 Super Users / System Administrators (Tier 1)",
         "• Onboarding: Add new members in the Members Directory to trigger automated activation links.\n"
         "• Maintenance: Review the Email Queue for delivery failures, check audit logs, and download weekly encrypted backups.\n"
         "• Verification: Use the Persona Switcher to inspect the user interface from the perspective of any active account.")
    ]
    
    for title, desc in workflows:
        doc.add_heading(title, level=2)
        doc.add_paragraph(desc)
        
    doc.add_heading("5. Troubleshooting & Frequently Asked Questions", level=1)
    faqs = [
        ("Sync & Multi-Device Updates", "The application automatically polls and syncs data every 7 seconds. To force an immediate refresh, press F5 / Ctrl+R or navigate to any sidebar module."),
        ("Expired Activation Links", "Activation tokens expire after 48 hours for security. Contact an administrator to click 'Resend Activation Link' from the Members Directory."),
        ("QR Scanner Camera Permissions", "Ensure camera access is allowed in your browser settings. In low-light environments, tap the Flashlight button in the scanner or enter the serial manually."),
        ("Receipt Upload Failures", "Ensure the receipt scan is in JPG, PNG, or PDF format and does not exceed 10 MB."),
        ("Apple Wallet Installation", "Attendees must open the pass link in Safari on iOS 14+. If opened inside third-party apps (WhatsApp, Chrome), tap the browser icon to open in Safari.")
    ]
    for q, a in faqs:
        p = doc.add_paragraph()
        r_q = p.add_run(f"Q: {q}\n")
        r_q.bold = True
        r_q.font.color.rgb = RGBColor(15, 60, 110)
        p.add_run(f"A: {a}")
        
    doc.add_paragraph()
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("— End of LEADS ERP Operations Manual —\n© 2026 LEADS Next Gen Centre, M.S. Ramaiah University of Applied Sciences. All rights reserved.")
    r_end.font.size = Pt(9)
    r_end.font.italic = True
    r_end.font.color.rgb = RGBColor(120, 120, 120)
    
    os.makedirs("docs", exist_ok=True)
    target_path = os.path.join(os.getcwd(), "docs", "LEADS_ERP_User_Manual.docx")
    doc.save(target_path)
    print(f"Successfully generated manual at: {target_path}")

if __name__ == "__main__":
    create_manual()
